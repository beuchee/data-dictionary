"""Tabular files (CSV / TSV / TXT / XLSX) → header + rows of text. Twin of table.ts."""

import io
import math
from dataclasses import dataclass
from datetime import date, datetime, time

from app.services.mapping_import.errors import MappingImportError
from app.services.mapping_import.text import clean, decode, detect_delimiter, parse_delimited

TEXT_EXTENSIONS = (".csv", ".tsv", ".txt", ".tab")


@dataclass
class InputFile:
    name: str
    data: bytes


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    sheet_names: list[str] | None = None
    delimiter: str | None = None


def extension_of(name: str) -> str:
    base = name.rsplit("/", 1)[-1]
    dot = base.rfind(".")
    return base[dot:].lower() if dot > 0 else ""


def is_xlsx(file: InputFile) -> bool:
    return extension_of(file.name) == ".xlsx"


def is_tabular(file: InputFile) -> bool:
    ext = extension_of(file.name)
    return ext in (".xlsx", "") or ext in TEXT_EXTENSIONS


def _js_number(value: float) -> str:
    """A number as JS's String() prints it (the TS twin reads XLSX numbers so)."""
    if math.isfinite(value) and value == int(value) and abs(value) < 1e21:
        return str(int(value))
    text = repr(value)
    if "e" in text:
        mantissa, exp = text.split("e")
        sign = "-" if exp.startswith("-") else "+"
        text = f"{mantissa.removesuffix('.0')}e{sign}{exp.lstrip('+-').lstrip('0') or '0'}"
    return text


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _js_number(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.replace(microsecond=0, tzinfo=None).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def read_xlsx(data: bytes, sheet: str | None = None) -> tuple[list[list[str]], list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise MappingImportError("unsupported_file", "unreadable workbook") from None
    try:
        names = list(wb.sheetnames)
        name = sheet if sheet is not None else (names[0] if names else None)
        if name is None or name not in names:
            raise MappingImportError("unknown_sheet", sheet)
        cells = [[cell_text(v) for v in row] for row in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()
    # Trailing blank cells are padding of the sheet's dimension, not data.
    out = []
    for row in cells:
        while row and row[-1] == "":
            row.pop()
        if any(clean(v) != "" for v in row):
            out.append(row)
    return out, names


def read_cells(file: InputFile, options: dict) -> tuple[list[list[str]], list[str] | None, str | None]:
    """The physical rows of a tabular file (no header handling)."""
    if not is_tabular(file):
        raise MappingImportError("unsupported_file", file.name)
    if is_xlsx(file):
        cells, names = read_xlsx(file.data, options.get("sheet"))
        return cells, names, None
    text = decode(file.data, options.get("encoding") or "auto")
    delimiter = options.get("delimiter") or "auto"
    if delimiter == "auto":
        delimiter = detect_delimiter(text)
    return parse_delimited(text, delimiter), None, delimiter


def table_from_cells(cells: list[list[str]], options: dict) -> Table:
    """Headers are trimmed; an empty or repeated one is renamed `column_<n>`."""
    rows = cells[max(0, int(options.get("skipRows") or 0)) :]
    if not rows:
        raise MappingImportError("empty_file")
    width = max(len(r) for r in rows)
    has_header = options.get("hasHeader", True)
    if has_header is None:
        has_header = True
    seen: set[str] = set()
    headers = []
    for i in range(width):
        h = clean(rows[0][i]) if has_header and i < len(rows[0]) else ""
        if not h or h in seen:
            h = f"column_{i + 1}"
        seen.add(h)
        headers.append(h)
    return Table(headers=headers, rows=rows[1:] if has_header else rows)


def read_table(file: InputFile, options: dict) -> Table:
    cells, names, delimiter = read_cells(file, options)
    table = table_from_cells(cells, options)
    table.sheet_names = names
    table.delimiter = delimiter
    return table
